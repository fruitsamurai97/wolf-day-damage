# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.chart.data import CategoryChartData
p = Presentation('Damage_Intelligence_EN_v2.pptx')

# M1 — translate chart categories (slide index 3)
for sh in p.slides[3].shapes:
    if sh.has_chart:
        cd = CategoryChartData()
        cd.categories = ['Parking', 'Manoeuvring', 'Rear-end collision', 'Stone chip', 'Hail', 'Vandalism', 'Wildlife']
        cd.add_series('', (110, 76, 62, 55, 55, 48, 44))
        sh.chart.replace_data(cd)

# helper to set a run's text preserving format
def replace_in_shape(shape, mapping):
    if not shape.has_text_frame: return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            for old, new in mapping.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)

sl = p.slides[4]  # results slide (agent "Slide 5")
shapes = list(sl.shapes)
# M4 — 99,9 % -> 99.9%
for sh in shapes:
    replace_in_shape(sh, {'99,9 %': '99.9%', '99,9%': '99.9%'})
# M5 — footnote "full 977" self-contradiction -> 434 damage cases
for sh in shapes:
    replace_in_shape(sh, {
        'On the full 977: DeepSeek 68.7% vs a 79.5% data-ceiling.': 'On the 434 damage cases: DeepSeek severity 68.7% vs a 79.5% data-ceiling.',
        'On the full 977': 'On the 434 damage cases',
    })

p.save('Damage_Intelligence_EN_v2.pptx')
print('deck fixed')
# verify
p2 = Presentation('Damage_Intelligence_EN_v2.pptx')
for sh in p2.slides[3].shapes:
    if sh.has_chart:
        print('new cats:', list(sh.chart.plots[0].categories))
