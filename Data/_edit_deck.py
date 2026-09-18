# -*- coding: utf-8 -*-
from pptx import Presentation
import copy
p = Presentation('Presentation_Hackathon_Damage_Intelligence.pptx')

def set_para(par, text):
    if par.runs:
        par.runs[0].text = text
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.text = text

def setshape(shape, lines):
    tf = shape.text_frame
    paras = tf.paragraphs
    for i, line in enumerate(lines):
        if i < len(paras):
            set_para(paras[i], line)
        else:
            src = paras[0]
            newp = copy.deepcopy(src._p)
            src._p.addnext(newp)
            paras = tf.paragraphs
            set_para(paras[i], line)
    for j in range(len(lines), len(tf.paragraphs)):
        set_para(tf.paragraphs[j], '')

S = p.slides
def sh(si, i): return list(S[si].shapes)[i]

# ---- SLIDE 0 ----
setshape(sh(0,2), ['Reading messy German workshop notes,', 'one case at a time'])
setshape(sh(0,3), ['A visual damage sheet per vehicle,', 'and an operational read on 1,000 dealership cases'])
setshape(sh(0,4), ['1,000 synthetic cases · German kept in German'])
setshape(sh(0,5), ['Wolf Day · DaiL · Track B'])
S[0].notes_slide.notes_text_frame.text = (
 "Opening pitch: \"Before we dive in, here's the live link so you can follow along: "
 "wolf-day-damage.vercel.app/dashboard. I'll walk you through what we built, then we'll "
 "review the dashboard together \u2014 it's designed for the workshop mechanic.\"")

# ---- SLIDE 1 (pipeline) ----
setshape(sh(1,0), ['One extraction feeds both deliverables'])
setshape(sh(1,3), ['German note'])
setshape(sh(1,4), ['DeepSeek reads one note and fills a strict form: parts, side, damage, severity \u2014 with a verbatim quote for each.'])
setshape(sh(1,6), ['Validated JSON'])
setshape(sh(1,7), ['A Pydantic schema forces a closed vocabulary. Unmapped terms are flagged, never dropped.'])
setshape(sh(1,9), ['Repair rules'])
setshape(sh(1,10), ['Replace, repair and remove-and-install come from deterministic code \u2014 no LLM guessing.'])
setshape(sh(1,12), ['Visualisation + analysis'])
setshape(sh(1,13), ['The same record colours the vehicle sheet and powers the corpus statistics.'])
setshape(sh(1,14), ['Case status is already structured (213 workshop statuses). Batch is idempotent: per-case cache + retry.'])

# ---- SLIDE 2 (case) ----
setshape(sh(2,0), ['A damage, its proof, and its operations'])
setshape(sh(2,3), ['Case 820002 · original note (kept in German)'])
setshape(sh(2,5), ['Windscreen: replace'])
setshape(sh(2,6), ['Action explicitly stated in the note'])
setshape(sh(2,7), ['Trim / plate: remove-and-install'])
setshape(sh(2,8), ['Derived from the adjacency table (not the note)'])
setshape(sh(2,9), ['Prototype capture. Colours separate damage from remove-and-install.'])

# ---- SLIDE 3 (corpus) ----
setshape(sh(3,0), ['What the corpus already shows'])
setshape(sh(3,2), ['Damage kinds · 450 cases'])
setshape(sh(3,4), ['41.3%'])
setshape(sh(3,5), ['of damage is parking', 'or manoeuvring', '(186 of 450 cases)'])
setshape(sh(3,6), ['668 open cases'])
setshape(sh(3,7), ['Median age: 150 days', 'A queue to review with the shop'])
setshape(sh(3,8), ['Synthetic data. Age since creation is not a repair-time measure.'])

# ---- SLIDE 4 (results) ----
setshape(sh(4,0), ['Measured on held-out cases'])
setshape(sh(4,2), ['DeepSeek extraction'])
setshape(sh(4,3), ['977 cases, outside examples and pilot'])
setshape(sh(4,5), ['Damage-zone F1'])
setshape(sh(4,7), ['Case kind + insurance', '100% quotes verbatim · 0 cases lost'])
setshape(sh(4,8), ['Severity: 5-model shootout on 150 cases'])
setshape(sh(4,9), ['Same test set for every model'])
setshape(sh(4,11), ['\u00b9 Insurance: 434 damage cases. Severity: 284 train / 150 test. On the full 977: DeepSeek 68.7% vs a 79.5% data-ceiling.'])
# table
tbl = [x for x in S[4].shapes if x.has_table][0].table
tr = {'Mod\u00e8le':'Model','Exactitude':'Accuracy','GBERT ajust\u00e9':'gbert (fine-tuned)',
      'R\u00e9gression logistique':'Logistic Regression'}
for row in tbl.rows:
    for cell in row.cells:
        txt = cell.text
        if txt in tr: cell.text = tr[txt]
        elif ',' in txt and '%' in txt: cell.text = txt.replace(',', '.')

# ---- SLIDE 5 (limits) — R&I limitation explicit ----
setshape(sh(5,0), ['Limits \u2014 said out loud'])
setshape(sh(5,2), ['The question the jury pressed on'])
setshape(sh(5,3), ['"How do we know which parts to remove to repair another one?" Today that comes from a hand-reviewed adjacency table (bodywork only).'])
setshape(sh(5,4), ['The dataset has NO parts, operation or removal labels \u2014 so remove-and-install is expert-rule-based and unverified. Mechanical R&I (brake pads \u2192 wheels) is out of scope.'])
setshape(sh(5,5), ['Severity is a data limit: the identical sentence carries light, medium AND severe labels \u2014 a ~79.5% ceiling, not a model gap.'])
setshape(sh(5,6), ['Next'])
setshape(sh(5,7), ['Validate the adjacency', 'rules with the workshop'])
setshape(sh(5,8), ['Learn R&I from real', 'repair-order histories'])
setshape(sh(5,9), ['Test on real', 'production notes'])
setshape(sh(5,10), ['Synthetic corpus: a regex baseline looks strong on zones; only 2 unmapped terms in 1,000 cases, flagged never dropped.'])

p.save('Damage_Intelligence_EN_v2.pptx')
print('saved Damage_Intelligence_EN_v2.pptx')
