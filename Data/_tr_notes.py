# -*- coding: utf-8 -*-
from pptx import Presentation
p = Presentation('Damage_Intelligence_EN_v2.pptx')

NOTES = {
1: '''50 seconds. The AI fills a strict schema and must keep a German excerpt justifying each damage. The closed vocabulary links a part and its side to the visualisation. Statuses come from structured data. An explicit action in the note takes priority; otherwise, rules propose the operations. Remove-and-install (R&I) uses a dependency table — it separates a damaged part from a neighbouring part that must be taken off. The per-case cache avoids redoing successful calls. Failures are flagged; calls retry with increasing back-off. Passing format validation does not prove the meaning is correct.

Local sources:
Work/pipeline/extract.py
Work/pipeline/schema.py
Work/pipeline/repair_logic.py
Work/pipeline/build_dataset.py''',
2: '''45 seconds. Show the German sentence before the colours. The note says the windscreen cannot be repaired and needs replacing. The windscreen shows orange, which here matches the extracted medium severity. Blue marks a trim piece to remove and reinstall per the business rule, not a second damage. This example is from the development pilot: it illustrates how it works, not proof of generalisation. The prototype also has a 3D view, but the capture uses the 2D view, clearer for explaining the data-to-zone link. Local demo: /dashboard/cases/820002/.

Local sources:
Data/frontend/public/data/cases.json
Data/frontend/src/sections/damage/car-diagram.tsx
Work/pipeline/repair_logic.py
Work/pipeline/ml_common.py''',
3: '''40 seconds. The corpus has 550 service cases and 450 damage cases. Among the latter, 110 are parking and 76 are manoeuvring, i.e. 186/450 = 41.3%. A business angle would be to look at their bodywork load, but we do not measure an operational gain here. The 668 cases on the open worklist have a median age of 150 days as of 18 September 2026. That is time since creation, not repair time, and some statuses may need checking. The observations describe this synthetic set, with no extrapolation to the market. The source notes stay in German.

Local sources:
Data/frontend/public/data/aggregates.json
Work/pipeline/build_dataset.py
Dataset/DATASET.md''',
4: '''60 seconds. Two protocols appear here. On the left, the DeepSeek extraction is evaluated on 977 cases, excluding the 3 prompt examples and the 20 pilot cases. Zones give 751 true positives, zero false positives and one false negative — an F1 of 99.93%. Case kind reaches 100%. Insurance reaches 100% on the 434 damage cases of this test. DeepSeek severity on those 434 cases is 68.7%, not to be compared directly with the table on the right. On the right, every model uses exactly the same 150 held-out ids for severity. The supervised models have 284 training examples. Verified counts: XGBoost 123/150, gbert 118/150, logistic regression 110/150, DeepSeek 109/150, LightGBM 98/150. These results describe a single split of the synthetic corpus. Recurring phrasings limit their reach. The lexical baseline also reaches 100% F1 on the zones of this corpus — so the scores do not prove a universal superiority of the LLM.

Local sources:
Work/pipeline/eval_report.json
Work/pipeline/ml_split.json
Work/pipeline/gbert_result.json
Data/frontend/public/data/model_comparison.json
Work/pipeline/train_cpu_models.py''',
5: '''45 seconds. Be explicit about the limits: the corpus is synthetic and reuses phrasings. Test ids are separate, but close phrasings can appear on both sides of the split. Validation by phrasing family, then on real notes, would be more convincing. The report shows the same stone-chip sentence can carry several severities. It estimates a 79.5% score for a text-fingerprint match on the full damage set: that estimate is specific to its method and scope, not a universal bound for the 150-case test (where the same method gives ~88.7%). We do not oppose it to XGBoost's 82%. Verbatim quotes and Pydantic enable auditing but do not guarantee semantic correctness. The R&I and replacement rules have no ground-truth labels in this set. Next steps: measure latency and real cost, plus external validation on real notes. No financial figure is claimed.

Local sources:
docs/jury-reliability-report.md
Work/pipeline/repair_logic.py
Dataset/DATASET.md
Work/pipeline/eval_report.json''',
}
for idx, text in NOTES.items():
    p.slides[idx].notes_slide.notes_text_frame.text = text
p.save('Damage_Intelligence_EN_v2.pptx')
print('notes translated to English')
# verify no French left in notes
p2=Presentation('Damage_Intelligence_EN_v2.pptx')
import re
bad=0
for i,s in enumerate(p2.slides):
    if s.has_notes_slide:
        t=s.notes_slide.notes_text_frame.text
        for w in ['Durée','dossiers','gravité','répar','français','pièce','manœuvre','plafond']:
            if w in t: print(f'slide {i}: still has \"{w}\"'); bad+=1
print('French leftovers in notes:', bad)
